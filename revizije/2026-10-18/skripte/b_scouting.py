# -*- coding: utf-8 -*-
"""Revizija 18.10.2026, drugi dio — scouting top 150: dokazi i izmjene (18.10.2026 12:47).

Korisnikova odluka 2a: Excel i Supabase mijenjati SAMO kad postojeći podaci daju dovoljno snažan
razlog; razdvojiti (1) statistički dokaz, (2) rezultate mečeva i (3) procjenu stila; sačuvati
staru verziju, zapisati razlog svake promjene i omogućiti povratak.

Pravila (PRAGOVI_B_2026-10-18.md, točka S, 12:20 — zapisana prije ovog računa):
  statistika = TML, 365 dana do 23.09.2026, glavni ždrijeb, >= 25 mečeva sa statistikom;
  "veliki server / servis je oružje" OPOVRGNUT ako su asovi po servis-gemu < medijana I hold < medijana;
  "odličan / jak povratnik"          OPOVRGNUT ako je povrat < medijana;
  "slab povrat"                      OPOVRGNUT ako je povrat > 75. percentila;
  "slab servis"                      OPOVRGNUT ako je hold > 75. percentila.
  Rezultati (Tokio/Peking/Šangaj) se zapisuju ZASEBNO i sami NE mijenjaju profil.
  Razine pouzdanosti se NE mijenjaju. Rang (činjenica, kod ga ne čita) se osvježava iz ATP singl
  liste koju dnevni run bilježi od 05.10.2026.

Izmjena teksta: originalna rečenica OSTAJE, a uz nju se doda oznaka "[18.10.2026 check: ...]"
s brojkama — model tako vidi i staru tvrdnju i da je izmjereno opovrgnuta.

    python revizije/2026-10-18/skripte/b_scouting.py               # suho: samo prijedlog
    python revizije/2026-10-18/skripte/b_scouting.py --primijeni   # kopija + upis + Excel
Povratak: python revizije/2026-10-18/skripte/b_scouting_vrati.py
"""
import json
import os
import re
import shutil
import sys
import unicodedata
from collections import defaultdict

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
sys.path.insert(0, ROOT)
from dotenv import load_dotenv  # noqa: E402
load_dotenv(os.path.join(ROOT, ".env"))
from database import supabase_client as db  # noqa: E402

OUTD = os.path.join(ROOT, "revizije", "2026-10-18", "scouting")
XLSX = os.path.join(ROOT, "ATP_Player_Scouting top150.xlsx")
DUMP = os.path.join(ROOT, ".cache", "rev_2026-10-18b")
TML_END = pd.Timestamp("2026-09-23")
MIN_MATCHES = 25
STAMP = "18.10.2026"
EN = {"veliki server": "big server", "jak povratnik": "strong returner", "slab povrat": "weak return",
      "slab servis": "weak serve"}

CLAIMS = {   # (oznaka, polja u kojima tvrdnja smije stajati, obrazac)
    "veliki server": (("style", "strengths"),
                      r"\bbig[- ]serv|\bhuge serve|\bmonster serve|\bpower server|\bserve[- ]first|\bbig-ace serve|\bstrong serve\b"),
    "jak povratnik": (("style", "strengths"),
                      r"\breturner\b|\belite return|\bstrong return\b|\bexcellent return|\bgreat return"),
    "slab povrat": (("weaknesses",),
                    r"return (is the weaker half|game is the weak link)|\b(weak|modest|poor|limited) return|return weaker"),
    "slab servis": (("weaknesses",),
                    r"serve is a liability|\bweak serve|\bsoft serve|serve not yet a weapon"),
}


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().replace("-", " ").replace("'", " ").strip().split())


def tml_profiles():
    tml = pd.read_pickle(os.path.join(ROOT, ".cache", "lab", "tml.pkl"))
    d = pd.to_datetime(tml["approx_date"])
    t = tml[(tml["src"] == "main") & (d > TML_END - pd.Timedelta(days=365)) & (d <= TML_END)]
    acc = defaultdict(lambda: defaultdict(float))
    names = {}
    cols = ["svpt", "ace", "1stIn", "1stWon", "2ndWon", "SvGms", "bpSaved", "bpFaced"]
    for r in t.itertuples(index=False):
        w = [getattr(r, f"w_{c}") for c in cols]
        l_ = [getattr(r, f"l_{c}") for c in cols]
        if any(pd.isna(x) for x in w + l_) or w[0] <= 0 or l_[0] <= 0 or w[5] <= 0 or l_[5] <= 0:
            continue
        for pid, nm, me, op, surf in ((r.winner_id, r.winner_name, w, l_, r.surface),
                                      (r.loser_id, r.loser_name, l_, w, r.surface)):
            names[pid] = nm
            a = acc[pid]
            a["n"] += 1
            for c, v in zip(cols, me):
                a[c] += v
            a["o_svpt"] += op[0]
            a["o_won"] += op[3] + op[4]
            if surf == "Hard":
                a["n_hard"] += 1
            elif surf == "Clay":
                a["n_clay"] += 1
    prof = {}
    for pid, a in acc.items():
        if a["n"] < MIN_MATCHES:
            continue
        prof[norm(names[pid])] = {
            "n": int(a["n"]), "name": names[pid],
            "ace_sg": a["ace"] / a["SvGms"], "hold": 1 - (a["bpFaced"] - a["bpSaved"]) / a["SvGms"],
            "ret": 1 - a["o_won"] / a["o_svpt"], "srv": (a["1stWon"] + a["2ndWon"]) / a["svpt"],
            "w1": a["1stWon"] / a["1stIn"], "w2": a["2ndWon"] / (a["svpt"] - a["1stIn"]),
        }
    return prof


def pct(prof, k, q):
    return float(pd.Series([p[k] for p in prof.values()]).quantile(q))


def find(prof, name):
    k = norm(name)
    if k in prof:
        return prof[k]
    parts = k.split()
    cand = [v for kk, v in prof.items() if parts and kk.split()[-1:] == parts[-1:]
            and kk.split()[0][:1] == parts[0][:1]]
    return cand[0] if len(cand) == 1 else None


def recent_ranks():
    am = json.load(open(os.path.join(DUMP, "analyzed_matches.json"), encoding="utf-8"))
    out = {}
    for r in sorted(am, key=lambda x: str(x.get("match_date"))):
        if str(r.get("match_date"))[:10] < "2026-10-05":
            continue
        cs = r.get("context_snapshot") or {}
        for side in ("p1", "p2"):
            rk = cs.get(f"{side}_ranking")
            if isinstance(rk, (int, float)) and 0 < rk < 999:
                out[norm(r.get("player1" if side == "p1" else "player2"))] = (int(rk), str(r["match_date"])[:10])
    return out


def asia_results():
    am = json.load(open(os.path.join(DUMP, "analyzed_matches.json"), encoding="utf-8"))
    order = {"R128": 1, "R64": 2, "R32": 3, "R16": 4, "QF": 5, "SF": 6, "F": 7}
    res = defaultdict(lambda: defaultdict(lambda: {"w": 0, "l": 0, "best": ""}))
    for r in am:
        t = str(r.get("tournament") or "")
        ev = ("Tokio" if "Japan Open" in t else "Peking" if "China Open" in t
              else "Šangaj" if "Shanghai" in t else None)
        aw = r.get("actual_winner")
        if not ev or aw not in (r.get("player1"), r.get("player2")):
            continue
        rd = (r.get("round") or "").upper()
        for p in (r.get("player1"), r.get("player2")):
            x = res[norm(p)][ev]
            if aw == p:
                x["w"] += 1
            else:
                x["l"] += 1
            if order.get(rd, 0) > order.get(x["best"], 0):
                x["best"] = rd
    return res


# NAKNADNO 18.10.2026 12:52 (PRAGOVI_B): tvrdnja je opovrgnuta tek uz odstupanje >= 2pp (hold,
# povrat) odnosno >= 0,05 asova po gemu od granice. Izvorno pravilo (bez razmaka) opovrglo je i
# Đokovića za 0,9pp (unutar ~1 standardne pogreške, bez prilagodbe protivnicima, uz RapidAPI koji
# kaže suprotno) — to nije "dovoljno snažan razlog". Stroži uvjet može samo smanjiti broj izmjena.
MARGIN_PP = 0.02
MARGIN_ACE = 0.05
BACKUP = os.path.join(OUTD, "player_scouting_prije_2026-10-18.json")
WITHDRAWN = [{"player": "Novak Djokovic", "polje": "style",
              "zapis_12_50": "[18.10.2026 check: 'strong returner' NOT supported - return points won 34.7% "
                             "(median 35.6%) ...]",
              "razlog": "povučeno 12:52: 0,9pp ispod medijana je unutar ~1 standardne pogreške (28 mečeva), "
                        "bez prilagodbe jačini protivnika; RapidAPI (06.09.) kaže 87. percentil. Rang ostaje."}]


def main(apply=False):
    live = db._select("player_scouting", select="*", limit=400)
    print(f"profila u Supabaseu: {len(live)}")
    # Izračun UVIJEK od stanja PRIJE 18.10. (sigurnosna kopija se nikad ne prepisuje), pa je skripta
    # ponovljiva: upisuje samo razliku između željenog i trenutnog stanja (i time povlači izmjene
    # koje više ne prolaze pravilo).
    if os.path.exists(BACKUP):
        rows = json.load(open(BACKUP, encoding="utf-8"))
        print(f"polazno stanje: sigurnosna kopija ({len(rows)} redaka)")
    else:
        rows = live
    prof = tml_profiles()
    med = {k: pct(prof, k, 0.5) for k in ("ace_sg", "hold", "ret")}
    p75 = {k: pct(prof, k, 0.75) for k in ("hold", "ret")}
    print(f"TML 365 d do {TML_END.date()}, glavni ždrijeb, >= {MIN_MATCHES} mečeva: {len(prof)} igrača | "
          f"medijan asovi/gem {med['ace_sg']:.3f}, hold {100 * med['hold']:.1f}%, povrat {100 * med['ret']:.1f}% | "
          f"75. pct hold {100 * p75['hold']:.1f}%, povrat {100 * p75['ret']:.1f}%")
    ranks, asia = recent_ranks(), asia_results()
    changes, evidence = [], []
    for r in rows:
        nm, disp = r["player_name"], r.get("display_name") or r["player_name"]
        p = find(prof, disp) or find(prof, nm)
        ev = {"player_name": nm, "display_name": disp, "confidence": r.get("confidence"),
              "rank_old": r.get("rank"), "stat": p, "claims": [], "asia": dict(asia.get(norm(disp), {}))}
        new_fields = {}
        for claim, (fields, pat) in CLAIMS.items():
            for fld in fields:
                txt = r.get(fld) or ""
                if not re.search(pat, txt, re.I):
                    continue
                if p is None:
                    ev["claims"].append((claim, fld, "neprovjerljivo (manje od 25 mečeva glavnog ždrijeba)"))
                    continue
                if claim == "veliki server":
                    bad = p["ace_sg"] <= med["ace_sg"] - MARGIN_ACE and p["hold"] <= med["hold"] - MARGIN_PP
                    num = (f"aces/service game {p['ace_sg']:.2f} (median {med['ace_sg']:.2f}), hold "
                           f"{100 * p['hold']:.1f}% (median {100 * med['hold']:.1f}%)")
                elif claim == "jak povratnik":
                    bad = p["ret"] <= med["ret"] - MARGIN_PP
                    num = f"return points won {100 * p['ret']:.1f}% (median {100 * med['ret']:.1f}%)"
                elif claim == "slab povrat":
                    bad = p["ret"] >= p75["ret"] + MARGIN_PP
                    num = f"return points won {100 * p['ret']:.1f}% (75th pct {100 * p75['ret']:.1f}%)"
                else:
                    bad = p["hold"] >= p75["hold"] + MARGIN_PP
                    num = f"hold {100 * p['hold']:.1f}% (75th pct {100 * p75['hold']:.1f}%)"
                ev["claims"].append((claim, fld, ("OPOVRGNUTO: " if bad else "nije opovrgnuto: ") + num))
                if bad and f"[{STAMP} check" not in txt:
                    base = new_fields.get(fld, txt)
                    new_fields[fld] = (base + f" [{STAMP} check: '{EN[claim]}' NOT supported - {num}; "
                                       f"ATP main draw, {p['n']} matches to 23.09.2026]")
        rk = ranks.get(norm(disp)) or ranks.get(norm(nm))
        if rk and rk[0] != r.get("rank"):
            new_fields["rank"] = rk[0]
            ev["rank_new"] = rk
        if new_fields:
            reasons = [c for c in ev["claims"] if c[2].startswith("OPOVRGNUTO")]
            changes.append({"player_name": nm, "display_name": disp, "confidence": r.get("confidence"),
                            "old": {k: r.get(k) for k in new_fields}, "new": new_fields,
                            "reasons": [f"{c[0]} ({c[1]}): {c[2]}" for c in reasons]
                            + ([f"rang: ATP singl lista iz dnevnog runa {rk[1]}"] if "rank" in new_fields else [])})
        evidence.append(ev)
    n_txt = sum(1 for c in changes if set(c["new"]) - {"rank"})
    n_rank = sum(1 for c in changes if "rank" in c["new"])
    print(f"\nprijedlog: {len(changes)} profila — tekst stila {n_txt}, rang {n_rank}")
    from collections import Counter
    cnt = Counter((c[0], c[2].split(":")[0]) for e in evidence for c in e["claims"])
    for k, v in sorted(cnt.items()):
        print(f"  {k[0]:14s} {k[1]:55s} {v}")
    for c in changes:
        if set(c["new"]) - {"rank"}:
            print(f"\n  * {c['display_name']} [{c['confidence']}]")
            for x in c["reasons"]:
                print(f"      {x}")
    if not apply:
        print("\n(suho — ništa nije upisano; za upis: --primijeni)")
        return changes, evidence
    os.makedirs(OUTD, exist_ok=True)
    if not os.path.exists(BACKUP):
        json.dump(rows, open(BACKUP, "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
        shutil.copy2(XLSX, os.path.join(OUTD, "ATP_Player_Scouting_top150_prije_2026-10-18.xlsx"))
        print(f"\nkopija: {BACKUP} ({len(rows)} redaka) + Excel")
    want = {r["player_name"]: dict(r) for r in rows}
    for c in changes:
        want[c["player_name"]].update(c["new"])
    cur = {r["player_name"]: r for r in live}
    ok = n_diff = 0
    for nm, w in want.items():
        diff = {f: w.get(f) for f in ("rank", "style", "strengths", "weaknesses")
                if cur.get(nm, {}).get(f) != w.get(f)}
        if not diff:
            continue
        n_diff += 1
        res = db._update("player_scouting", diff, {"player_name": f"eq.{nm}"})
        ok += 1 if isinstance(res, list) and res else 0
    print(f"upisano u Supabase (razlika prema trenutnom stanju): {ok}/{n_diff}")
    json.dump({"vrijeme": "18.10.2026 12:50, ispravak 12:52",
               "pravila": "PRAGOVI_B_2026-10-18.md, točka S (12:20) i NAKNADNO 12:52",
               "promjene": changes, "povuceno": WITHDRAWN},
              open(os.path.join(OUTD, "promjene_2026-10-18.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1, default=str)
    write_excel(evidence, changes, med, p75)
    return changes, evidence


def write_excel(evidence, changes, med, p75):
    """Glavni list = isti izvoz kao `scripts/export_scouting_excel.py`; uz njega dva nova lista."""
    sys.path.insert(0, os.path.join(ROOT, "scripts"))
    import export_scouting_excel as ex
    ex.main()
    import openpyxl
    from openpyxl.styles import Font, Alignment
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.create_sheet("Dokazi 18.10.2026")
    head = ["Igrač", "Pouzdanost profila",
            "STATISTIKA (TML 365 d, glavni ždrijeb): mečeva", "asovi / servis-gem", "hold %",
            "povrat %", "servis %", "1. servis %", "2. servis %",
            "REZULTATI (naš uzorak): Tokio", "Peking", "Šangaj",
            "PROCJENA STILA: tvrdnje iz profila i provjera"]
    for c, h in enumerate(head, 1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
    ws.cell(row=2, column=1, value=(f"Medijan: asovi/gem {med['ace_sg']:.2f}, hold {100 * med['hold']:.1f}%, povrat "
                                    f"{100 * med['ret']:.1f}% | 75. percentil: hold {100 * p75['hold']:.1f}%, povrat "
                                    f"{100 * p75['ret']:.1f}%. Rezultati su OPIS i sami ne mijenjaju profil."))
    for i, e in enumerate(sorted(evidence, key=lambda x: (x.get("rank_old") or 999)), start=3):
        p = e["stat"] or {}
        a = e["asia"]
        fmt = lambda ev: (f"{a[ev]['w']}-{a[ev]['l']} ({a[ev]['best']})" if ev in a else "")
        vals = [e["display_name"], e["confidence"], p.get("n"),
                round(p["ace_sg"], 2) if p else None, round(100 * p["hold"], 1) if p else None,
                round(100 * p["ret"], 1) if p else None, round(100 * p["srv"], 1) if p else None,
                round(100 * p["w1"], 1) if p else None, round(100 * p["w2"], 1) if p else None,
                fmt("Tokio"), fmt("Peking"), fmt("Šangaj"),
                " | ".join(f"{c[0]} ({c[1]}): {c[2]}" for c in e["claims"])]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v).alignment = Alignment(vertical="top", wrap_text=(c == 13))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["M"].width = 110
    ws2 = wb.create_sheet("Promjene 18.10.2026")
    for c, h in enumerate(["Igrač", "Polje", "Staro", "Novo", "Razlog"], 1):
        ws2.cell(row=1, column=c, value=h).font = Font(bold=True)
    k = 2
    for ch in changes:
        for fld, new in ch["new"].items():
            for c, v in enumerate([ch["display_name"], fld, ch["old"].get(fld), new, "; ".join(ch["reasons"])], 1):
                ws2.cell(row=k, column=c, value=v).alignment = Alignment(vertical="top", wrap_text=True)
            k += 1
    for wd in WITHDRAWN:
        for c, v in enumerate([wd["player"], wd["polje"], wd["zapis_12_50"], "(vraćeno na staro)", wd["razlog"]], 1):
            ws2.cell(row=k, column=c, value=v).alignment = Alignment(vertical="top", wrap_text=True)
        k += 1
    for col, w in zip("ABCDE", (26, 12, 60, 60, 80)):
        ws2.column_dimensions[col].width = w
    wb.save(XLSX)
    print(f"Excel osvježen: {os.path.basename(XLSX)} (+ listovi 'Dokazi 18.10.2026' i 'Promjene 18.10.2026')")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(apply="--primijeni" in sys.argv)
